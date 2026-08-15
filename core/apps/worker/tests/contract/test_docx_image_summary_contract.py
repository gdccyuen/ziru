from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path
from typing import Any

os.environ.setdefault("DATABASE_URL", "postgresql+asyncpg://test:test@localhost/test")
os.environ.setdefault("TMP_PATH", "/tmp/knowhere-test")
os.environ.setdefault("S3_BUCKET_NAME", "test-uploads")
os.environ.setdefault("S3_ACCESS_KEY_ID", "test")
os.environ.setdefault("S3_SECRET_ACCESS_KEY", "test")
os.environ.setdefault("S3_TEMP_PATH", "/tmp")

import gevent  # noqa: E402
from docx import Document  # noqa: E402
from PIL import Image  # noqa: E402
from pytest import MonkeyPatch  # noqa: E402

from app.services.document_parser.formats.docx.asset_store import (  # noqa: E402
    DocxAssetStore,
)
from app.services.document_parser.formats.docx.image_summary_scheduler import (  # noqa: E402
    DocxImageSummaryScheduler,
)
from app.services.document_parser.formats.docx.parser import (  # noqa: E402
    handle_image,
    handle_table,
)
from shared.core.config import settings  # noqa: E402
from shared.services.ai.summary.model import AssetSummary  # noqa: E402


def test_docx_inline_image_summaries_are_bounded(
    tmp_path: Path,
    monkeypatch: MonkeyPatch,
) -> None:
    active_calls = 0
    max_active_calls = 0
    call_paths: list[str] = []

    def fake_summarize(**kwargs: Any) -> AssetSummary:
        nonlocal active_calls, max_active_calls
        active_calls += 1
        max_active_calls = max(max_active_calls, active_calls)
        call_paths.append(kwargs["image_paths"][0])
        gevent.sleep(0.01)
        active_calls -= 1
        return AssetSummary(summary=f"summary {len(call_paths)}", title="title")

    monkeypatch.setattr(settings, "DOCX_IMAGE_SUMMARY_MAX_CONCURRENT", 4)
    monkeypatch.setattr("shared.services.ai.summary.engine.summarize", fake_summarize)

    asset_store = _create_asset_store(tmp_path)
    headings_stack = _create_headings_stack()
    rows: list[list[object]] = []
    seen_images: dict[str, dict[str, str]] = {}
    scheduler = DocxImageSummaryScheduler(should_summarize=True)

    for image_index in range(8):
        _, rows, is_new_image = handle_image(
            rows,
            _create_image_meta(image_index),
            asset_store,
            headings_stack,
            "Heading",
            image_index,
            smart_summary=True,
            seen_images=seen_images,
            image_summary_scheduler=scheduler,
        )
        assert is_new_image is True

    scheduler.run_all()

    assert len(call_paths) == 8
    assert max_active_calls <= 4
    assert all("summary" in str(row[0]) for row in rows)


def test_docx_inline_image_title_remains_in_searchable_path(
    tmp_path: Path,
    monkeypatch: MonkeyPatch,
) -> None:
    def fake_summarize(**_: Any) -> AssetSummary:
        return AssetSummary(summary="revenue chart summary", title="Revenue Chart")

    monkeypatch.setattr(settings, "DOCX_IMAGE_SUMMARY_MAX_CONCURRENT", 4)
    monkeypatch.setattr("shared.services.ai.summary.engine.summarize", fake_summarize)

    asset_store = _create_asset_store(tmp_path)
    headings_stack = _create_headings_stack()
    rows: list[list[object]] = []
    seen_images: dict[str, dict[str, str]] = {}
    scheduler = DocxImageSummaryScheduler(should_summarize=True)

    _, rows, is_new_image = handle_image(
        rows,
        _create_image_meta(1),
        asset_store,
        headings_stack,
        "Heading",
        0,
        smart_summary=True,
        seen_images=seen_images,
        image_summary_scheduler=scheduler,
    )
    scheduler.run_all()

    relative_path = str(rows[0][1])
    assert is_new_image is True
    assert relative_path == "images/image-1 Revenue Chart.png"
    assert "Revenue Chart" in str(rows[0][0])
    assert rows[0][12] == "Revenue Chart"
    assert (tmp_path / relative_path).exists()


def test_docx_duplicate_images_share_one_summary_call(
    tmp_path: Path,
    monkeypatch: MonkeyPatch,
) -> None:
    call_paths: list[str] = []

    def fake_summarize(**kwargs: Any) -> AssetSummary:
        call_paths.append(kwargs["image_paths"][0])
        return AssetSummary(summary="shared duplicate summary", title="shared title")

    monkeypatch.setattr(settings, "DOCX_IMAGE_SUMMARY_MAX_CONCURRENT", 4)
    monkeypatch.setattr("shared.services.ai.summary.engine.summarize", fake_summarize)

    asset_store = _create_asset_store(tmp_path)
    headings_stack = _create_headings_stack()
    rows: list[list[object]] = []
    seen_images: dict[str, dict[str, str]] = {}
    scheduler = DocxImageSummaryScheduler(should_summarize=True)
    image_meta = _create_image_meta(1)

    _, rows, first_is_new = handle_image(
        rows,
        image_meta,
        asset_store,
        headings_stack,
        "Heading",
        0,
        smart_summary=True,
        seen_images=seen_images,
        image_summary_scheduler=scheduler,
    )
    _, rows, second_is_new = handle_image(
        rows,
        image_meta,
        asset_store,
        headings_stack,
        "Heading",
        1,
        smart_summary=True,
        seen_images=seen_images,
        image_summary_scheduler=scheduler,
    )

    scheduler.run_all()

    assert first_is_new is True
    assert second_is_new is False
    assert len(call_paths) == 1
    assert len(rows) == 2
    assert all("shared duplicate summary" in str(row[0]) for row in rows)
    assert all("shared title" in str(row[1]) for row in rows)
    assert all("shared duplicate summary" in str(item) for item in headings_stack[-1]["content"][1:])


def test_docx_summary_image_false_makes_no_image_summary_calls(
    tmp_path: Path,
    monkeypatch: MonkeyPatch,
) -> None:
    calls: list[str] = []

    def fake_summarize(**kwargs: Any) -> AssetSummary:
        calls.append(kwargs["image_paths"][0])
        return AssetSummary(summary="unexpected")

    monkeypatch.setattr("shared.services.ai.summary.engine.summarize", fake_summarize)

    asset_store = _create_asset_store(tmp_path)
    headings_stack = _create_headings_stack()
    rows: list[list[object]] = []
    seen_images: dict[str, dict[str, str]] = {}
    scheduler = DocxImageSummaryScheduler(should_summarize=False)

    _, rows, is_new_image = handle_image(
        rows,
        _create_image_meta(1),
        asset_store,
        headings_stack,
        "Heading",
        0,
        smart_summary=False,
        seen_images=seen_images,
        image_summary_scheduler=scheduler,
    )
    scheduler.run_all()

    assert is_new_image is True
    assert calls == []
    assert rows[0][5] == "image-1"
    assert "unexpected" not in str(rows[0][0])


def test_docx_table_image_summary_is_applied_before_html_render(
    tmp_path: Path,
    monkeypatch: MonkeyPatch,
) -> None:
    calls: list[str] = []

    def fake_summarize(**kwargs: Any) -> AssetSummary:
        calls.append(kwargs["usage_task"])
        return AssetSummary(summary="table cell summary", title="table image title")

    monkeypatch.setattr(settings, "DOCX_IMAGE_SUMMARY_MAX_CONCURRENT", 4)
    monkeypatch.setattr("shared.services.ai.summary.engine.summarize", fake_summarize)

    asset_store = _create_asset_store(tmp_path)
    headings_stack = _create_headings_stack()
    rows: list[list[object]] = []
    seen_images: dict[str, dict[str, str]] = {}
    scheduler = DocxImageSummaryScheduler(should_summarize=True)
    table = _create_docx_table()

    _, rows, next_image_count = handle_table(
        rows,
        table,
        asset_store,
        headings_stack,
        "Heading",
        0,
        summary_table=False,
        summary_image=True,
        cell_images={(0, 1): [_create_image_meta(1)]},
        img_count=0,
        seen_images=seen_images,
        image_summary_scheduler=scheduler,
    )

    assert calls == ["parser.docx.table_image"]
    assert next_image_count == 1
    assert rows[0][2] == "image"
    assert rows[0][5] == "image-1\ntable cell summary"
    assert rows[0][12] == "table image title"
    assert "table image title" not in str(rows[0][1])
    assert rows[1][2] == "table"

    table_path = tmp_path / str(rows[1][1])
    table_html = table_path.read_text(encoding="utf-8")
    assert "<td>first</td><td>second<br/><em>[table cell summary]</em></td>" in table_html


def test_docx_image_summary_failure_keeps_fallback_reference(
    tmp_path: Path,
    monkeypatch: MonkeyPatch,
) -> None:
    def fake_summarize(**_: Any) -> AssetSummary:
        raise RuntimeError("provider failed")

    monkeypatch.setattr(settings, "DOCX_IMAGE_SUMMARY_MAX_CONCURRENT", 4)
    monkeypatch.setattr("shared.services.ai.summary.engine.summarize", fake_summarize)

    asset_store = _create_asset_store(tmp_path)
    headings_stack = _create_headings_stack()
    rows: list[list[object]] = []
    seen_images: dict[str, dict[str, str]] = {}
    scheduler = DocxImageSummaryScheduler(should_summarize=True)

    _, rows, is_new_image = handle_image(
        rows,
        _create_image_meta(1),
        asset_store,
        headings_stack,
        "Heading",
        0,
        smart_summary=True,
        seen_images=seen_images,
        image_summary_scheduler=scheduler,
    )
    scheduler.run_all()

    assert is_new_image is True
    assert rows[0][5] == "image-1"
    assert rows[0][12] == ""
    assert str(rows[0][0]).startswith("\n[images/")
    assert "provider failed" not in str(rows[0][0])


def _create_asset_store(tmp_path: Path) -> DocxAssetStore:
    asset_store = DocxAssetStore(str(tmp_path))
    asset_store.reset()
    return asset_store


def _create_headings_stack() -> list[dict[str, Any]]:
    return [{"level": 1, "heading": "Heading", "content": ["nearby context"]}]


def _create_image_meta(seed: int) -> dict[str, Any]:
    return {
        "image_name": f"image-{seed}.png",
        "data": _create_image_bytes(seed),
        "size": 12 * 1024,
    }


def _create_image_bytes(seed: int) -> bytes:
    image = Image.new(
        "RGB",
        (32, 32),
        (seed % 255, (seed * 23) % 255, (seed * 47) % 255),
    )
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _create_docx_table() -> Any:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "first"
    table.cell(0, 1).text = "second"
    return table
